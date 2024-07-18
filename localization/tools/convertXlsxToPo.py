import os
import sys
import traceback
import openpyxl
import polib

def _convertXlsxToPo(xlsxFilePath, poFilePath):
    workbook = openpyxl.load_workbook(xlsxFilePath)
    worksheet = workbook.active
    po = polib.POFile()
    for row in range(2, worksheet.max_row + 1):
        msgid = worksheet.cell(row=row, column=1).value
        msgstr = worksheet.cell(row=row, column=2).value
        if msgid is None:
            continue
        if msgstr is None:
            msgstr = ""
        entry = polib.POEntry(
            msgid=str(msgid),
            msgstr=str(msgstr),
        )
        po.append(entry)
    po.save(poFilePath)
        
if __name__ == '__main__':
    try:
        xlsxFilePath = sys.argv[1]
        poFilePath = os.path.splitext(os.path.basename(xlsxFilePath))[0] + ".po";
        _convertXlsxToPo(xlsxFilePath, poFilePath)
    except:
        traceback.print_exc()
        os.system("pause")
        
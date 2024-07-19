import os
import sys
import traceback
import openpyxl
import polib
from openpyxl.styles import Alignment

def _preprocessPo(srcPoFile, tempPoFile):
    rows = []
    with open(srcPoFile, 'r', encoding="utf-8") as reader:
        for line in reader.readlines():
            if line.find('#~') == 0:
                continue
            rows.append(line)
    with open(tempPoFile, 'w', encoding="utf-8") as writer:
        for line in rows:
            writer.write(line)

def convertPoToXlsx(poPilePath, xlsxFilePath):
    _preprocessPo(poPilePath, "temp.po")

    po = polib.pofile("temp.po")
    workbook = openpyxl.Workbook()
    worksheet = workbook.active

    worksheet.column_dimensions['A'].width = 30
    worksheet.column_dimensions['B'].width = 50
    worksheet.column_dimensions['C'].width = 50

    worksheet['A1'] = 'msgid'
    worksheet['B1'] = 'msgstr'
    worksheet['C1'] = 'comments'

    for i, entry in enumerate(po, start=2):
        worksheet.cell(row=i, column=1, value=entry.msgid)
        worksheet.cell(row=i, column=2, value=entry.msgstr)
        worksheet.cell(row=i, column=3, value=entry.comment)

    alignment = Alignment(horizontal='center', vertical='center')    
    for row in range(1, len(po) + 2):
        for col in range(1, 4):
            worksheet.cell(row=row, column=col).alignment = alignment

    workbook.save(xlsxFilePath)
        
if __name__ == '__main__':
    try:
        xlsxFilePath = sys.argv[1]
        poFilePath = os.path.splitext(os.path.basename(xlsxFilePath))[0] + ".po";
        convertPoToXlsx(xlsxFilePath, poFilePath)
    except:
        traceback.print_exc()
        os.system("pause")
        